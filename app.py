"""
Web Tabanlı B2B Pazarlama ve Harita Kontrol Paneli (Flask Uygulaması)
Ankara İncek / LÖSANTE Karşısı Ticari Mülk Sistemi
Erişim: http://localhost:5000
"""

import os
import sys
import json
import webbrowser
import threading
from flask import Flask, render_template, request, jsonify, send_file, send_from_directory

# Force UTF-8
sys.stdout.reconfigure(encoding="utf-8")

# Add current dir to path
APP_DIR = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, APP_DIR)

from modules.enrichment import clean_phone_for_whatsapp, generate_whatsapp_url
from modules.map_generator import generate_interactive_map
from modules.pitch_generator import generate_pitch_document
from modules.scoring import calculate_propfit_score
from modules.decision_maker_hunter import find_decision_makers, generate_confidential_teaser
from modules.mailer import (
    load_smtp_config, save_smtp_config, test_smtp_connection,
    build_email_content, send_proposal_email, load_sent_logs,
    find_proposal_file, is_gmail_api_ready
)
from google_maps_lead_harvester import scrape_google_maps_places, export_to_excel, load_config

import openpyxl

app = Flask(__name__, template_folder="templates")

# File Paths - Supports both workspace parent dir and standalone clone
LOCAL_DATA_DIR = os.path.join(APP_DIR, "data")
PARENT_DIR = os.path.abspath(os.path.join(APP_DIR, ".."))

if os.path.exists(os.path.join(PARENT_DIR, "GOOGLE_HARITA_TOPLANAN_KURUMLAR.xlsx")):
    DATA_ROOT = PARENT_DIR
elif os.path.exists(os.path.join(LOCAL_DATA_DIR, "GOOGLE_HARITA_TOPLANAN_KURUMLAR.xlsx")):
    DATA_ROOT = LOCAL_DATA_DIR
else:
    DATA_ROOT = PARENT_DIR

MASTER_EXCEL_PATH = os.path.join(DATA_ROOT, "GOOGLE_HARITA_TOPLANAN_KURUMLAR.xlsx")
STRATEGIC_EXCEL_PATH = os.path.join(DATA_ROOT, "STRATEJIK_PAZARLAMA_VE_HEDEF_MUSTERI_LISTESI.xlsx")
MAP_HTML_PATH = os.path.join(DATA_ROOT, "INCEK_TICARI_HEDEF_HARITASI.html")
PROPOSALS_DIR = os.path.join(DATA_ROOT, "TEKLIF_MEKTUPLARI")


def load_leads_from_excel():
    """Master Excel dosyasından toplanan kurumları okur ve zenginleştirir."""
    leads = []
    if not os.path.exists(MASTER_EXCEL_PATH):
        return leads
        
    try:
        wb = openpyxl.load_workbook(MASTER_EXCEL_PATH, data_only=True)
        if "TOPLANAN_KURUMLAR" in wb.sheetnames:
            ws = wb["TOPLANAN_KURUMLAR"]
            for r in range(2, ws.max_row + 1):
                name = ws.cell(row=r, column=2).value
                if not name:
                    continue
                phone = str(ws.cell(row=r, column=7).value or "N/A")
                wa_phone = clean_phone_for_whatsapp(phone)
                
                dist_val = ws.cell(row=r, column=11).value
                try:
                    dist_float = float(dist_val) if dist_val is not None else None
                except Exception:
                    dist_float = None
                    
                lead_obj = {
                    "kurum_adi": name,
                    "ana_kategori": ws.cell(row=r, column=3).value or "",
                    "is_kolu": ws.cell(row=r, column=4).value or "Kurumsal",
                    "adres": ws.cell(row=r, column=5).value or "N/A",
                    "plus_kodu": ws.cell(row=r, column=6).value or "N/A",
                    "telefon": phone,
                    "telefon_wa": wa_phone or "",
                    "eposta": ws.cell(row=r, column=8).value or "N/A",
                    "web_sitesi": ws.cell(row=r, column=9).value or "",
                    "koordinatlar": ws.cell(row=r, column=10).value or "N/A",
                    "mesafe_km": dist_float,
                    "puan": ws.cell(row=r, column=12).value or "",
                    "hbu_model": ws.cell(row=r, column=13).value or "Ticari Yerleşke",
                    "oncelik": ws.cell(row=r, column=14).value or "B+",
                    "maps_url": ws.cell(row=r, column=15).value or "#",
                    "tarih": ws.cell(row=r, column=16).value or ""
                }
                scoring = calculate_propfit_score(lead_obj)
                lead_obj.update(scoring)
                leads.append(lead_obj)
                
        # Sort by PropFit score descending
        leads = sorted(leads, key=lambda x: x.get("propfit_score", 0), reverse=True)
        for idx, item in enumerate(leads):
            item["index"] = idx
    except Exception as e:
        print(f"Excel okunurken hata: {e}")
        
    return leads

def compute_stats(leads):
    """Kurumlar üzerinden özet analitik metrikleri hesaplar."""
    total = len(leads)
    health_count = sum(1 for l in leads if any(k in (l.get("is_kolu") or "").lower() for k in ["sağlık", "tıp", "onkoloji", "fizik"]))
    edu_count = sum(1 for l in leads if any(k in (l.get("is_kolu") or "").lower() for k in ["eğitim", "kolej", "okul"]))
    dental_count = sum(1 for l in leads if any(k in (l.get("is_kolu") or "").lower() for k in ["diş", "estetik", "cerrahi"]))
    aplus_count = sum(1 for l in leads if l.get("oncelik") == "A+")
    
    tier1_count = sum(1 for l in leads if "Tier 1" in l.get("tier", ""))
    tier2_count = sum(1 for l in leads if "Tier 2" in l.get("tier", ""))
    tier3_count = sum(1 for l in leads if "Tier 3" in l.get("tier", ""))
    
    valid_scores = [l.get("propfit_score", 0) for l in leads]
    avg_propfit = round(sum(valid_scores) / len(valid_scores), 1) if valid_scores else 0.0
    
    valid_dists = [l["mesafe_km"] for l in leads if l.get("mesafe_km") is not None]
    avg_dist = round(sum(valid_dists) / len(valid_dists), 1) if valid_dists else 0.0
    
    return {
        "total_leads": total,
        "health_count": health_count,
        "edu_count": edu_count,
        "dental_count": dental_count,
        "aplus_count": aplus_count,
        "tier1_count": tier1_count,
        "tier2_count": tier2_count,
        "tier3_count": tier3_count,
        "avg_propfit": avg_propfit,
        "avg_distance": avg_dist
    }

@app.route("/")
def dashboard():
    leads = load_leads_from_excel()
    stats = compute_stats(leads)
    email_logs = load_sent_logs()
    for l in leads:
        k_name = l.get("kurum_adi", "")
        if k_name in email_logs:
            l["email_sent"] = email_logs[k_name]
        else:
            l["email_sent"] = None
    smtp_cfg = load_smtp_config()
    safe_smtp_cfg = dict(smtp_cfg)
    safe_smtp_cfg["has_password"] = bool(smtp_cfg.get("smtp_password"))
    safe_smtp_cfg["smtp_password"] = "***" if safe_smtp_cfg["has_password"] else ""
    return render_template("dashboard.html", leads=leads, stats=stats, smtp_config=safe_smtp_cfg, email_logs=email_logs)

@app.route("/map")
def serve_map():
    if not os.path.exists(MAP_HTML_PATH):
        leads = load_leads_from_excel()
        generate_interactive_map(leads, MAP_HTML_PATH)
    return send_file(MAP_HTML_PATH)

@app.route("/download/master_excel")
def download_master_excel():
    if os.path.exists(MASTER_EXCEL_PATH):
        return send_file(MASTER_EXCEL_PATH, as_attachment=True)
    return "Dosya bulunamadı", 404

@app.route("/download/strategic_excel")
def download_strategic_excel():
    if os.path.exists(STRATEGIC_EXCEL_PATH):
        return send_file(STRATEGIC_EXCEL_PATH, as_attachment=True)
    return "Dosya bulunamadı", 404

@app.route("/api/generate_pitch/<int:lead_index>")
def api_generate_pitch(lead_index):
    leads = load_leads_from_excel()
    if 0 <= lead_index < len(leads):
        lead = leads[lead_index]
        file_path = generate_pitch_document(lead, PROPOSALS_DIR)
        return send_file(file_path, as_attachment=True)
    return jsonify({"error": "Kurum bulunamadı"}), 404

@app.route("/api/generate_all_pitches")
def api_generate_all_pitches():
    leads = load_leads_from_excel()
    count = 0
    for lead in leads:
        generate_pitch_document(lead, PROPOSALS_DIR)
        count += 1
    return jsonify({
        "status": "success",
        "message": f"Toplam {count} adet kurum için kişiselleştirilmiş Word (.docx) teklif mektubu '{PROPOSALS_DIR}' klasörüne başarıyla üretildi!"
    })

@app.route("/api/generate_teaser/<int:lead_index>")
def api_generate_teaser(lead_index):
    leads = load_leads_from_excel()
    if 0 <= lead_index < len(leads):
        lead = leads[lead_index]
        teaser_text = generate_confidential_teaser(lead)
        return jsonify({
            "status": "success",
            "kurum_adi": lead.get("kurum_adi"),
            "teaser": teaser_text
        })
    return jsonify({"error": "Kurum bulunamadı"}), 404

@app.route("/api/hunt_executives/<int:lead_index>")
def api_hunt_executives(lead_index):
    leads = load_leads_from_excel()
    if 0 <= lead_index < len(leads):
        lead = leads[lead_index]
        website = lead.get("web_sitesi")
        execs = find_decision_makers(website)
        return jsonify({
            "status": "success",
            "kurum_adi": lead.get("kurum_adi"),
            "executives": execs
        })
    return jsonify({"error": "Kurum bulunamadı"}), 404

@app.route("/api/email/preview/<int:lead_index>")
def api_email_preview(lead_index):
    leads = load_leads_from_excel()
    if 0 <= lead_index < len(leads):
        lead = leads[lead_index]
        content = build_email_content(lead)
        cfg = load_smtp_config()
        doc_path = find_proposal_file(lead.get("kurum_adi", ""), PROPOSALS_DIR)
        return jsonify({
            "status": "success",
            "kurum_adi": lead.get("kurum_adi"),
            "recipient_email": lead.get("eposta", "") if lead.get("eposta") != "N/A" else "",
            "sender_email": cfg.get("sender_email", "yigit.narin@cb.com.tr"),
            "sender_name": cfg.get("sender_name", "Yiğit Narin | Coldwell Banker VIP"),
            "subject": content["subject"],
            "html_body": content["html_body"],
            "text_body": content["text_body"],
            "attachment_file": os.path.basename(doc_path) if doc_path else None,
            "has_attachment": bool(doc_path and os.path.exists(doc_path)),
            "smtp_ready": bool(cfg.get("smtp_password")) or is_gmail_api_ready(),
            "gmail_api_ready": is_gmail_api_ready()
        })
    return jsonify({"error": "Kurum bulunamadı"}), 404

@app.route("/api/email/send/<int:lead_index>", methods=["POST"])
def api_email_send(lead_index):
    leads = load_leads_from_excel()
    if 0 <= lead_index < len(leads):
        lead = leads[lead_index]
        req_data = request.json or {}
        recipient = req_data.get("recipient_email") or lead.get("eposta")
        custom_subject = req_data.get("subject")
        custom_html = req_data.get("html_body")
        
        # Ensure proposal file exists, if not generate it
        doc_path = find_proposal_file(lead.get("kurum_adi", ""), PROPOSALS_DIR)
        if not doc_path:
            generate_pitch_document(lead, PROPOSALS_DIR)
            
        res = send_proposal_email(
            lead_data=lead,
            proposals_dir=PROPOSALS_DIR,
            recipient_email=recipient,
            custom_subject=custom_subject,
            custom_html=custom_html
        )
        return jsonify(res)
    return jsonify({"success": False, "message": "Kurum bulunamadı"}), 404

@app.route("/api/email/config", methods=["GET", "POST"])
def api_email_config():
    if request.method == "POST":
        data = request.json or {}
        updated = save_smtp_config(data)
        safe = dict(updated)
        safe["has_password"] = bool(safe.get("smtp_password"))
        safe["smtp_password"] = "***" if safe["has_password"] else ""
        return jsonify({"status": "success", "config": safe})
    else:
        cfg = load_smtp_config()
        safe = dict(cfg)
        safe["has_password"] = bool(safe.get("smtp_password"))
        safe["smtp_password"] = "***" if safe["has_password"] else ""
        return jsonify(safe)

@app.route("/api/email/test", methods=["POST"])
def api_email_test():
    data = request.json or {}
    res = test_smtp_connection(data if "smtp_password" in data else None)
    return jsonify(res)

@app.route("/api/email/logs")
def api_email_logs():
    logs = load_sent_logs()
    return jsonify(logs)


@app.route("/api/email/auth_google", methods=["POST"])
def api_email_auth_google():
    try:
        import subprocess
        script_path = os.path.join(APP_DIR, "google_auth_setup.py")
        subprocess.Popen([sys.executable, script_path], creationflags=subprocess.CREATE_NEW_CONSOLE)
        return jsonify({
            "status": "success",
            "message": "Yetkilendirme penceresi açıldı. Tarayıcınızda açılan Google onay ekranında yigit.narin@cb.com.tr hesabınızı seçip 'İzin Ver' deyiniz."
        })
    except Exception as e:
        return jsonify({"status": "error", "message": str(e)}), 500


@app.route("/api/scan", methods=["POST"])
def api_start_scan():
    data = request.json or {}
    category = data.get("category", "all")
    custom_query = data.get("query")
    max_places = int(data.get("max_places", 6))
    
    config = load_config()
    search_queries = []
    
    if custom_query and custom_query.strip():
        search_queries = [custom_query.strip()]
    else:
        presets = config.get("category_presets", {})
        if category == "all":
            for cat_key, cat_data in presets.items():
                search_queries.extend(cat_data.get("keywords", []))
        elif category in presets:
            search_queries = presets[category].get("keywords", [])
        else:
            search_queries = [f"{category} İncek Ankara"]
            
    # Run scraper in current context
    anchor_lat = config.get("anchor_property", {}).get("latitude", 39.8245)
    anchor_lon = config.get("anchor_property", {}).get("longitude", 32.7485)
    
    results = scrape_google_maps_places(
        search_queries=search_queries,
        anchor_lat=anchor_lat,
        anchor_lon=anchor_lon,
        max_places_per_query=max_places,
        headless=True
    )
    
    # Export to Excel
    export_to_excel(results, output_path=MASTER_EXCEL_PATH)
    
    # Regenerate Map
    all_leads = load_leads_from_excel()
    generate_interactive_map(all_leads, MAP_HTML_PATH)
    
    return jsonify({
        "status": "success",
        "new_count": len(results),
        "total_count": len(all_leads)
    })

def open_browser():
    time.sleep(1.5)
    webbrowser.open("http://localhost:5000")

if __name__ == "__main__":
    print("\n" + "="*75)
    print("  COLDWELL BANKER VIP - INCEK TICARI MULK B2B PAZARLAMA KONTROL PANELI")
    print("="*75)
    print("  Web Paneli Baslatiliyor: http://localhost:5000")
    print("  Durdurmak icin: CTRL + C tuslarina basiniz.")
    print("="*75 + "\n")
    
    threading.Thread(target=open_browser, daemon=True).start()
    app.run(host="0.0.0.0", port=5000, debug=False)
