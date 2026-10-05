"""
Google Maps Kurumsal Veri Toplama ve B2B Aday Otomasyonu
(Ankara İncek / LÖSANTE Karşısı Ticari Mülk Pazarlama Sistemi)

Bu otomasyon; belirlenen koordinat veya adres çevresinde, önerilen stratejik iş kollarındaki
(Sağlık/Medikal, Kolej/Eğitim, Savunma/Teknoloji, Hukuk, Diplomatik) kurumları Google Maps üzerinden
otomatik olarak tarar; Kurum Adı, Kategori, Adres, Plus Kodu, Telefon, Web Sitesi, E-Posta,
Koordinat ve Mülke Mesafeyi çekerek kurumsal formatta Excel (XLSX) dosyasına kaydeder.
"""

import os
import sys
import re
import json
import math
import time
import argparse
import urllib.parse
import warnings
from datetime import datetime

# Suppress urllib3 warnings
warnings.filterwarnings("ignore")

# Force UTF-8 output for Windows console
sys.stdout.reconfigure(encoding="utf-8")

try:
    import requests
    from bs4 import BeautifulSoup
    from playwright.sync_api import sync_playwright
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter
except ImportError as e:
    print(f"Eksik kütüphane: {e}")
    print("Lütfen gerekli kütüphaneleri yükleyin: pip install playwright openpyxl beautifulsoup4 requests")
    sys.exit(1)

# Base directories
SCRIPT_DIR = os.path.dirname(os.path.abspath(__file__))
CONFIG_PATH = os.path.join(SCRIPT_DIR, "config.json")

# Default Anchor: LÖSANTE / Kızılcaşar Mah. 2705. Cadde No: 23, Gölbaşı / Ankara
DEFAULT_ANCHOR_LAT = 39.8245
DEFAULT_ANCHOR_LON = 32.7485
DEFAULT_OUTPUT_XLSX = (
    os.path.abspath(os.path.join(SCRIPT_DIR, "..", "GOOGLE_HARITA_TOPLANAN_KURUMLAR.xlsx"))
    if os.path.exists(os.path.abspath(os.path.join(SCRIPT_DIR, "..", "GOOGLE_HARITA_TOPLANAN_KURUMLAR.xlsx")))
    else os.path.join(SCRIPT_DIR, "data", "GOOGLE_HARITA_TOPLANAN_KURUMLAR.xlsx")
)

def load_config():
    if os.path.exists(CONFIG_PATH):
        try:
            with open(CONFIG_PATH, "r", encoding="utf-8") as f:
                return json.load(f)
        except Exception as e:
            print(f"Uyarı: config.json okunamadı ({e}), varsayılanlar kullanılacak.")
    return {
        "anchor_property": {
            "title": "İncek LÖSANTE Karşısı Çift Ruhsatlı Ticari Mülk",
            "latitude": DEFAULT_ANCHOR_LAT,
            "longitude": DEFAULT_ANCHOR_LON
        },
        "output_excel_path": DEFAULT_OUTPUT_XLSX
    }

def clean_text(text):
    """Metin içindeki özel unicode ikonları ve boşlukları temizler."""
    if not text:
        return ""
    text = re.sub(r'[\ue000-\uf8ff]', '', text)
    return text.strip()

def calculate_distance(lat1, lon1, lat2, lon2):
    """Haversine formülü ile iki koordinat arası kuş uçuşu mesafeyi km cinsinden hesaplar."""
    if lat1 is None or lon1 is None or lat2 is None or lon2 is None:
        return None
    R = 6371.0 # Dünya yarıçapı km
    dlat = math.radians(lat2 - lat1)
    dlon = math.radians(lon2 - lon1)
    a = math.sin(dlat / 2)**2 + math.cos(math.radians(lat1)) * math.cos(math.radians(lat2)) * math.sin(dlon / 2)**2
    c = 2 * math.atan2(math.sqrt(a), math.sqrt(1 - a))
    return round(R * c, 2)

def extract_email_from_website(url):
    """Web sitesi ana sayfasından ve /iletisim linkinden kurumsal e-posta adresini ayıklar."""
    if not url or url == "N/A" or not url.startswith("http"):
        return "N/A"
    
    headers = {
        "User-Agent": "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/128.0.0.0 Safari/537.36"
    }
    
    found_emails = set()
    test_urls = [url]
    
    # Try home page
    try:
        r = requests.get(url, headers=headers, timeout=4, verify=False)
        if r.status_code == 200:
            # find mailto: links first
            soup = BeautifulSoup(r.text, "html.parser")
            for a in soup.find_all("a", href=True):
                href = a["href"]
                if "mailto:" in href:
                    clean_m = href.replace("mailto:", "").split("?")[0].strip()
                    if "@" in clean_m:
                        found_emails.add(clean_m)
            
            # regex search
            emails = re.findall(r'[a-zA-Z0-9_.+-]+@[a-zA-Z0-9-]+\.[a-zA-Z0-9-.]+', r.text)
            for e in emails:
                if not e.lower().endswith(('.png', '.jpg', '.jpeg', '.webp', '.svg', '.gif', '.css', '.js')):
                    found_emails.add(e)
                    
            # If nothing found, check contact page link
            if not found_emails:
                for a in soup.find_all("a", href=True):
                    txt = a.get_text().lower()
                    href = a["href"].lower()
                    if "iletisim" in txt or "contact" in txt or "iletisim" in href or "contact" in href:
                        contact_url = urllib.parse.urljoin(url, a["href"])
                        try:
                            cr = requests.get(contact_url, headers=headers, timeout=4, verify=False)
                            if cr.status_code == 200:
                                c_emails = re.findall(r'[a-zA-Z0-9_.+-]+@[a-zA-Z0-9-]+\.[a-zA-Z0-9-.]+', cr.text)
                                for ce in c_emails:
                                    if not ce.lower().endswith(('.png', '.jpg', '.jpeg', '.webp', '.svg', '.gif')):
                                        found_emails.add(ce)
                        except Exception:
                            pass
                        break
    except Exception:
        pass

    # Filter out common garbage
    valid = [e for e in found_emails if not e.startswith(("sentry", "wix", "example", "domain"))]
    return valid[0] if valid else "N/A"

def determine_hbu_and_priority(sector_key, category, name, distance_km):
    """Mülkümüzün niteliklerine göre kullanım modeli ve öncelik skoru belirler."""
    name_cat = (name + " " + category).lower()
    
    # Default assignments
    if "tıp" in name_cat or "poliklinik" in name_cat or "hastane" in name_cat or "sağlık" in name_cat:
        hbu = "Butik Tıp Merkezi / Cerrahi Poliklinik / İhtisas Merkezi"
        sector_name = "Sağlık & Medikal"
    elif "diş" in name_cat or "dent" in name_cat or "ağız" in name_cat:
        hbu = "Müstakil VIP Ağız ve Diş Sağlığı Merkezi"
        sector_name = "Ağız ve Diş Sağlığı"
    elif "estetik" in name_cat or "cerrahi" in name_cat or "güzellik" in name_cat or "dermatoloji" in name_cat:
        hbu = "VIP Estetik & Plastik Cerrahi / Saç Ekim Köşkü"
        sector_name = "Plastik & Estetik Cerrahi"
    elif "fizik tedavi" in name_cat or "rehabilitasyon" in name_cat:
        hbu = "Müstakil Bahçeli Fizik Tedavi ve Rehabilitasyon Merkezi"
        sector_name = "Fizik Tedavi & Rehabilitasyon"
    elif "kolej" in name_cat or "okul" in name_cat or "eğitim" in name_cat or "anaokulu" in name_cat:
        hbu = "Müstakil Butik Kolej / VIP Anaokulu / Özel Eğitim Kampüsü"
        sector_name = "Özel Eğitim & Kolejler"
    elif "savunma" in name_cat or "radar" in name_cat or "havacılık" in name_cat:
        hbu = "Müstakil Şirket Genel Merkezi & Savunma Protokol Köşkü"
        sector_name = "Savunma Sanayii"
    elif "yazılım" in name_cat or "bilişim" in name_cat or "teknoloji" in name_cat or "siber" in name_cat:
        hbu = "Müstakil AR-GE ve Yazılım Geliştirme Karargâhı"
        sector_name = "Bilişim & Yazılım"
    elif "hukuk" in name_cat or "avukat" in name_cat:
        hbu = "Prestijli Müstakil Hukuk Bürosu & Arabuluculuk Merkezi"
        sector_name = "Prestij Hukuk"
    elif "denetim" in name_cat or "ymm" in name_cat or "muhasebe" in name_cat:
        hbu = "YMM Bağımsız Denetim ve Danışmanlık Hizmet Binası"
        sector_name = "Bağımsız Denetim & YMM"
    elif "elçilik" in name_cat or "büyükelçilik" in name_cat or "ataşelik" in name_cat or "embassy" in name_cat:
        hbu = "Diplomatik Rezidans / Konsolosluk Hizmet Binası"
        sector_name = "Diplomatik Misyon"
    else:
        hbu = "Müstakil Kurumsal Şirket Genel Merkezi"
        sector_name = "Kurumsal Merkez"
        
    # Priority calculation based on distance and sector
    if distance_km is not None and distance_km <= 3.5:
        priority = "A+"
    elif distance_km is not None and distance_km <= 7.0:
        priority = "A"
    else:
        priority = "B+"
        
    # Health sector directly opposite Lösante gets bump
    if sector_name in ["Sağlık & Medikal", "Ağız ve Diş Sağlığı", "Fizik Tedavi & Rehabilitasyon"]:
        if distance_km is not None and distance_km <= 5.0:
            priority = "A+"
            
    return sector_name, hbu, priority

def scrape_google_maps_places(search_queries, anchor_lat=DEFAULT_ANCHOR_LAT, anchor_lon=DEFAULT_ANCHOR_LON, max_places_per_query=8, headless=True):
    """Playwright kullanarak Google Maps üzerinde arama yapar ve detaylı bilgileri çeker."""
    all_results = []
    seen_places = set()
    
    print(f"\n[🚀 BAŞLATILIYOR] Google Maps Kurumsal Arama Motoru")
    print(f"Referans Mülk Koordinatları: {anchor_lat}, {anchor_lon} (İncek / LÖSANTE)")
    print(f"Toplam Arama Terimi Sayısı: {len(search_queries)}")
    print(f"Tarayıcı Modu: {'Gizli (Headless)' if headless else 'Görsel (Headed)'}\n")
    
    with sync_playwright() as p:
        browser = p.chromium.launch(headless=headless)
        context = browser.new_context(
            locale="tr-TR",
            user_agent="Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/128.0.0.0 Safari/537.36"
        )
        page = context.new_page()
        
        for q_idx, query in enumerate(search_queries, start=1):
            print(f"\n────────────────────────────────────────────────────────────")
            print(f"[{q_idx}/{len(search_queries)}] 🔍 Arama Yapılıyor: '{query}'")
            print(f"────────────────────────────────────────────────────────────")
            
            maps_url = f"https://www.google.com/maps/search/{urllib.parse.quote(query)}/@{anchor_lat},{anchor_lon},14z?hl=tr"
            
            try:
                page.goto(maps_url, wait_until="domcontentloaded", timeout=25000)
                page.wait_for_timeout(3000)
                
                # Cookie consent
                try:
                    consent = page.locator("button[aria-label*='Kabul'], button:has-text('Tümünü kabul et')")
                    if consent.count() > 0:
                        consent.first.click()
                        page.wait_for_timeout(2000)
                except Exception:
                    pass
                
                # Feed scrolling
                feed = page.locator("div[role='feed']")
                if feed.count() > 0:
                    scroll_count = min(4, math.ceil(max_places_per_query / 3))
                    for _ in range(scroll_count):
                        feed.evaluate("el => el.scrollTop += 1200")
                        page.wait_for_timeout(1200)
                        
                # Extract place links
                links = page.locator("a[href*='/maps/place/']").all()
                print(f"Bulunan İşletme / Kurum Sayısı: {len(links)}")
                
                collected_for_query = 0
                for link in links:
                    if collected_for_query >= max_places_per_query:
                        break
                        
                    aria_name = link.get_attribute("aria-label")
                    href = link.get_attribute("href")
                    
                    if not aria_name or aria_name in seen_places:
                        continue
                        
                    print(f"  👉 İnceleniyor: {aria_name}...")
                    
                    try:
                        link.click()
                        page.wait_for_timeout(2500)
                    except Exception:
                        continue
                        
                    detail_url = page.url
                    
                    # Extract Coordinates
                    coords_match = re.search(r'!3d([0-9\.]+)!4d([0-9\.]+)', detail_url)
                    if not coords_match:
                        coords_match = re.search(r'@([0-9\.]+),([0-9\.]+)', detail_url)
                        
                    lat, lon = None, None
                    distance_km = None
                    if coords_match:
                        lat, lon = float(coords_match.group(1)), float(coords_match.group(2))
                        distance_km = calculate_distance(anchor_lat, anchor_lon, lat, lon)
                        
                    # Extract Name
                    name = clean_text(aria_name)
                    
                    # Extract Category
                    cat_el = page.locator("button[jsaction*='category']")
                    category = clean_text(cat_el.first.inner_text()) if cat_el.count() > 0 else "Kurum"
                    
                    # Extract Address
                    addr_el = page.locator("button[data-item-id='address'], [data-tooltip*='Adres'], button[aria-label*='Adres:']")
                    address = clean_text(addr_el.first.inner_text()) if addr_el.count() > 0 else "N/A"
                    
                    # Extract Plus Code
                    plus_el = page.locator("button[data-item-id='oloc'], [data-tooltip*='Artı kod'], button[aria-label*='Artı kod:']")
                    plus_code = clean_text(plus_el.first.inner_text()) if plus_el.count() > 0 else "N/A"
                    
                    # Extract Phone
                    phone_el = page.locator("button[data-item-id*='phone'], [data-tooltip*='Telefon'], button[aria-label*='Telefon:']")
                    phone = clean_text(phone_el.first.inner_text()) if phone_el.count() > 0 else "N/A"
                    
                    # Extract Website
                    web_el = page.locator("a[data-item-id='authority'], [data-tooltip*='Web sitesi'], a[aria-label*='Web sitesi:']")
                    website = web_el.first.get_attribute("href") if web_el.count() > 0 else "N/A"
                    
                    # Extract Rating
                    rating_el = page.locator("div.fontDisplayLarge, span[aria-label*='yıldız']")
                    rating = clean_text(rating_el.first.inner_text()) if rating_el.count() > 0 else "N/A"
                    
                    # Extract Email from website if available
                    email = "N/A"
                    if website != "N/A":
                        email = extract_email_from_website(website)
                        
                    # Determine Sektör, HBU and Priority
                    sector_name, hbu_model, priority = determine_hbu_and_priority(query, category, name, distance_km)
                    
                    record = {
                        "kurum_adi": name,
                        "ana_kategori": category,
                        "is_kolu": sector_name,
                        "adres": address,
                        "plus_kodu": plus_code,
                        "telefon": phone,
                        "eposta": email,
                        "web_sitesi": website,
                        "koordinatlar": f"{lat}, {lon}" if lat and lon else "N/A",
                        "mesafe_km": distance_km,
                        "puan": rating,
                        "maps_url": detail_url,
                        "hbu_model": hbu_model,
                        "oncelik": priority,
                        "tarih": datetime.now().strftime("%d.%m.%Y %H:%M")
                    }
                    
                    all_results.append(record)
                    seen_places.add(name)
                    collected_for_query += 1
                    
                    dist_str = f"{distance_km} km" if distance_km else "Mesafe Yok"
                    print(f"     ✅ Eklendi: [{sector_name} | {priority}] {name} ({dist_str}) | Tel: {phone} | Mail: {email}")
                    
            except Exception as e:
                print(f"Hata oluştu ({query}): {e}")
                continue
                
        browser.close()
        
    print(f"\n[🏁 TAMAMLANDI] Toplam {len(all_results)} adet benzersiz kurum verisi başarıyla çekildi.")
    return all_results

def export_to_excel(results, output_path=DEFAULT_OUTPUT_XLSX):
    """Toplanan verileri kurumsal Excel (XLSX) formatında kaydeder veya günceller."""
    os.makedirs(os.path.dirname(os.path.abspath(output_path)), exist_ok=True)
    
    # Check existing data to merge and deduplicate
    existing_data = []
    seen_keys = set()
    
    if os.path.exists(output_path):
        try:
            wb_exist = openpyxl.load_workbook(output_path, data_only=True)
            if "TOPLANAN_KURUMLAR" in wb_exist.sheetnames:
                ws_e = wb_exist["TOPLANAN_KURUMLAR"]
                for r in range(2, ws_e.max_row + 1):
                    k_name = ws_e.cell(row=r, column=2).value
                    if k_name:
                        seen_keys.add(str(k_name).strip().lower())
                        existing_data.append({
                            "kurum_adi": ws_e.cell(row=r, column=2).value,
                            "ana_kategori": ws_e.cell(row=r, column=3).value,
                            "is_kolu": ws_e.cell(row=r, column=4).value,
                            "adres": ws_e.cell(row=r, column=5).value,
                            "plus_kodu": ws_e.cell(row=r, column=6).value,
                            "telefon": ws_e.cell(row=r, column=7).value,
                            "eposta": ws_e.cell(row=r, column=8).value,
                            "web_sitesi": ws_e.cell(row=r, column=9).value,
                            "koordinatlar": ws_e.cell(row=r, column=10).value,
                            "mesafe_km": ws_e.cell(row=r, column=11).value,
                            "puan": ws_e.cell(row=r, column=12).value,
                            "hbu_model": ws_e.cell(row=r, column=13).value,
                            "oncelik": ws_e.cell(row=r, column=14).value,
                            "maps_url": ws_e.cell(row=r, column=15).value,
                            "tarih": ws_e.cell(row=r, column=16).value
                        })
        except Exception as e:
            print(f"Mevcut dosya okunurken hata: {e}")

    # Append new results that haven't been seen
    new_added = 0
    for res in results:
        key = str(res["kurum_adi"]).strip().lower()
        if key not in seen_keys:
            existing_data.append(res)
            seen_keys.add(key)
            new_added += 1

    wb = openpyxl.Workbook()
    wb.remove(wb.active) # remove default sheet
    
    # -----------------------------------------------------------------
    # Styles
    # -----------------------------------------------------------------
    font_tbl_header = Font(name="Segoe UI", size=9, bold=True, color="FFFFFF")
    font_data = Font(name="Segoe UI", size=9, color="1E293B")
    font_data_bold = Font(name="Segoe UI", size=9, bold=True, color="1E293B")
    font_title = Font(name="Segoe UI", size=14, bold=True, color="FFFFFF")
    font_subtitle = Font(name="Segoe UI", size=10, italic=True, color="E2E8F0")
    
    fill_navy = PatternFill(start_color="1B365D", end_color="1B365D", fill_type="solid")
    fill_sec = PatternFill(start_color="2C3E50", end_color="2C3E50", fill_type="solid")
    fill_zebra = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")
    fill_white = PatternFill(start_color="FFFFFF", end_color="FFFFFF", fill_type="solid")
    
    fill_aplus = PatternFill(start_color="FEE2E2", end_color="FEE2E2", fill_type="solid")
    font_aplus = Font(name="Segoe UI", size=9, bold=True, color="991B1B")
    fill_a = PatternFill(start_color="FEF3C7", end_color="FEF3C7", fill_type="solid")
    font_a = Font(name="Segoe UI", size=9, bold=True, color="92400E")
    fill_bplus = PatternFill(start_color="E0F2FE", end_color="E0F2FE", fill_type="solid")
    font_bplus = Font(name="Segoe UI", size=9, bold=True, color="075985")
    
    thin_border = Border(
        left=Side(style="thin", color="D1D5DB"),
        right=Side(style="thin", color="D1D5DB"),
        top=Side(style="thin", color="D1D5DB"),
        bottom=Side(style="thin", color="D1D5DB")
    )
    
    # -----------------------------------------------------------------
    # Sheet 1: TOPLANAN_KURUMLAR (Master Table)
    # -----------------------------------------------------------------
    ws1 = wb.create_sheet(title="TOPLANAN_KURUMLAR")
    ws1.views.sheetView[0].showGridLines = True
    
    headers = [
        "Sıra",
        "Kurum / İşletme Adı",
        "Google Kategorisi",
        "Önerilen İş Kolu / Sektör",
        "Tam Adres Bilgisi",
        "Google Plus Kodu",
        "Telefon Numarası",
        "Kurumsal E-Posta",
        "Web Sitesi",
        "Koordinatlar (Enlem, Boylam)",
        "Mülke Mesafe (km)",
        "Google Puanı",
        "Önerilen HBU Kullanım Modeli",
        "Öncelik Derecesi",
        "Google Maps Bağlantısı",
        "Veri Çekilme Tarihi"
    ]
    
    ws1.row_dimensions[1].height = 28
    for c_idx, h_text in enumerate(headers, start=1):
        cell = ws1.cell(row=1, column=c_idx, value=h_text)
        cell.font = font_tbl_header
        cell.fill = fill_navy
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = thin_border
        
    for r_idx, item in enumerate(existing_data, start=2):
        ws1.row_dimensions[r_idx].height = 22
        is_zebra = (r_idx % 2 == 1)
        r_fill = fill_zebra if is_zebra else fill_white
        
        row_vals = [
            r_idx - 1,
            item.get("kurum_adi", ""),
            item.get("ana_kategori", ""),
            item.get("is_kolu", ""),
            item.get("adres", ""),
            item.get("plus_kodu", ""),
            item.get("telefon", ""),
            item.get("eposta", ""),
            item.get("web_sitesi", ""),
            item.get("koordinatlar", ""),
            item.get("mesafe_km", ""),
            item.get("puan", ""),
            item.get("hbu_model", ""),
            item.get("oncelik", ""),
            item.get("maps_url", ""),
            item.get("tarih", "")
        ]
        
        for c_idx, val in enumerate(row_vals, start=1):
            cell = ws1.cell(row=r_idx, column=c_idx, value=val)
            cell.font = font_data
            cell.fill = r_fill
            cell.border = thin_border
            
            if c_idx in [1, 6, 7, 10, 11, 12, 14, 16]:
                cell.alignment = Alignment(horizontal="center", vertical="center")
            elif c_idx in [5, 13]:
                cell.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
            else:
                cell.alignment = Alignment(horizontal="left", vertical="center")
                
            # Priority color pill
            if c_idx == 14: # Öncelik
                p_val = str(val).strip()
                if p_val == "A+":
                    cell.fill = fill_aplus
                    cell.font = font_aplus
                elif p_val == "A":
                    cell.fill = fill_a
                    cell.font = font_a
                elif p_val == "B+":
                    cell.fill = fill_bplus
                    cell.font = font_bplus

    # Set Column Widths for Sheet 1
    col_widths = {
        "A": 8,   # Sıra
        "B": 38,  # Kurum Adı
        "C": 28,  # Kategori
        "D": 26,  # İş Kolu
        "E": 45,  # Adres
        "F": 22,  # Plus Kodu
        "G": 18,  # Telefon
        "H": 28,  # E-Posta
        "I": 30,  # Web Sitesi
        "J": 24,  # Koordinatlar
        "K": 18,  # Mesafe
        "L": 14,  # Puan
        "M": 40,  # HBU Modeli
        "N": 14,  # Öncelik
        "O": 35,  # Maps URL
        "P": 18   # Tarih
    }
    for col_let, w in col_widths.items():
        ws1.column_dimensions[col_let].width = w

    # -----------------------------------------------------------------
    # Sheet 2: DASHBOARD_OZET
    # -----------------------------------------------------------------
    ws2 = wb.create_sheet(title="DASHBOARD_OZET")
    ws2.views.sheetView[0].showGridLines = True
    
    ws2.merge_cells("A1:F1")
    ws2["A1"] = "GOOGLE MAPS KURUMSAL VERİ TOPLAMA VE HEDEF KİTLE DASHBOARD"
    ws2["A1"].font = font_title
    ws2["A1"].fill = fill_navy
    ws2["A1"].alignment = Alignment(horizontal="center", vertical="center")
    ws2.row_dimensions[1].height = 36
    
    ws2.merge_cells("A2:F2")
    ws2["A2"] = f"Referans: İncek LÖSANTE Karşısı Ticari Mülk (Ada 119464 Parsel 13) | Toplam Kayıt: {len(existing_data)}"
    ws2["A2"].font = font_subtitle
    ws2["A2"].fill = fill_sec
    ws2["A2"].alignment = Alignment(horizontal="center", vertical="center")
    ws2.row_dimensions[2].height = 22
    
    # Calculate Sector Counts
    sector_counts = {}
    prio_counts = {"A+": 0, "A": 0, "B+": 0}
    for item in existing_data:
        s = item.get("is_kolu", "Diğer")
        sector_counts[s] = sector_counts.get(s, 0) + 1
        p = item.get("oncelik", "B+")
        if p in prio_counts:
            prio_counts[p] += 1
            
    # Table 1: Sector Breakdown
    ws2.merge_cells("A4:C4")
    ws2["A4"] = "SEKTÖREL DAĞILIM"
    ws2["A4"].font = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")
    ws2["A4"].fill = fill_sec
    ws2["A4"].alignment = Alignment(horizontal="left", vertical="center")
    
    ws2["A5"] = "İş Kolu / Sektör"
    ws2["A5"].font = font_tbl_header
    ws2["A5"].fill = fill_navy
    ws2["A5"].border = thin_border
    
    ws2["B5"] = "Kurum Sayısı"
    ws2["B5"].font = font_tbl_header
    ws2["B5"].fill = fill_navy
    ws2["B5"].border = thin_border
    ws2["B5"].alignment = Alignment(horizontal="center", vertical="center")
    
    ws2["C5"] = "Yüzde Payı"
    ws2["C5"].font = font_tbl_header
    ws2["C5"].fill = fill_navy
    ws2["C5"].border = thin_border
    ws2["C5"].alignment = Alignment(horizontal="center", vertical="center")
    ws2.row_dimensions[5].height = 22
    
    total_recs = max(1, len(existing_data))
    curr_r = 6
    for s_name, count in sorted(sector_counts.items(), key=lambda x: x[1], reverse=True):
        ws2[f"A{curr_r}"] = s_name
        ws2[f"A{curr_r}"].font = font_data
        ws2[f"A{curr_r}"].border = thin_border
        
        ws2[f"B{curr_r}"] = count
        ws2[f"B{curr_r}"].font = font_data_bold
        ws2[f"B{curr_r}"].border = thin_border
        ws2[f"B{curr_r}"].alignment = Alignment(horizontal="center", vertical="center")
        
        pct = f"%{round((count / total_recs) * 100, 1)}"
        ws2[f"C{curr_r}"] = pct
        ws2[f"C{curr_r}"].font = font_data
        ws2[f"C{curr_r}"].border = thin_border
        ws2[f"C{curr_r}"].alignment = Alignment(horizontal="center", vertical="center")
        ws2.row_dimensions[curr_r].height = 20
        curr_r += 1
        
    # Table 2: Priority Breakdown
    ws2.merge_cells("E4:F4")
    ws2["E4"] = "ÖNCELİK PUANLAMASI"
    ws2["E4"].font = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")
    ws2["E4"].fill = fill_sec
    ws2["E4"].alignment = Alignment(horizontal="left", vertical="center")
    
    ws2["E5"] = "Öncelik Derecesi"
    ws2["E5"].font = font_tbl_header
    ws2["E5"].fill = fill_navy
    ws2["E5"].border = thin_border
    
    ws2["F5"] = "Adet"
    ws2["F5"].font = font_tbl_header
    ws2["F5"].fill = fill_navy
    ws2["F5"].border = thin_border
    ws2["F5"].alignment = Alignment(horizontal="center", vertical="center")
    
    prio_r = 6
    for p_code in ["A+", "A", "B+"]:
        cnt = prio_counts.get(p_code, 0)
        ws2[f"E{prio_r}"] = f"{p_code} ({'En Yüksek' if p_code=='A+' else ('Yüksek' if p_code=='A' else 'Orta')})"
        ws2[f"E{prio_r}"].font = font_data_bold
        ws2[f"E{prio_r}"].border = thin_border
        if p_code == "A+":
            ws2[f"E{prio_r}"].fill = fill_aplus
            ws2[f"E{prio_r}"].font = font_aplus
        elif p_code == "A":
            ws2[f"E{prio_r}"].fill = fill_a
            ws2[f"E{prio_r}"].font = font_a
        elif p_code == "B+":
            ws2[f"E{prio_r}"].fill = fill_bplus
            ws2[f"E{prio_r}"].font = font_bplus
            
        ws2[f"F{prio_r}"] = cnt
        ws2[f"F{prio_r}"].font = font_data_bold
        ws2[f"F{prio_r}"].border = thin_border
        ws2[f"F{prio_r}"].alignment = Alignment(horizontal="center", vertical="center")
        ws2.row_dimensions[prio_r].height = 20
        prio_r += 1
        
    ws2.column_dimensions["A"].width = 35
    ws2.column_dimensions["B"].width = 16
    ws2.column_dimensions["C"].width = 16
    ws2.column_dimensions["D"].width = 6
    ws2.column_dimensions["E"].width = 25
    ws2.column_dimensions["F"].width = 16

    wb.save(output_path)
    print(f"\n[💾 KAYDEDİLDİ] Excel Dosyası: {output_path}")
    print(f"Toplam Kayıt: {len(existing_data)} (Yeni Eklenen: {new_added})")

def main():
    parser = argparse.ArgumentParser(description="Google Maps B2B Kurumsal Veri Toplama Otomasyonu")
    parser.add_argument("--category", "-c", default="all", choices=["all", "saglik", "egitim", "savunma_teknoloji", "hukuk_denetim", "diplomatik"],
                        help="Taranacak sektör kategorisi (varsayılan: all)")
    parser.add_argument("--query", "-q", default=None, help="Özel arama sorgusu (örn: 'diş polikliniği İncek Ankara')")
    parser.add_argument("--lat", type=float, default=DEFAULT_ANCHOR_LAT, help="Referans merkez enlem (Varsayılan: 39.8245 - LÖSANTE)")
    parser.add_argument("--lon", type=float, default=DEFAULT_ANCHOR_LON, help="Referans merkez boylam (Varsayılan: 32.7485 - LÖSANTE)")
    parser.add_argument("--max-places", "-m", type=int, default=6, help="Arama sorgusu başına çekilecek maksimum kurum sayısı (Varsayılan: 6)")
    parser.add_argument("--headed", action="store_true", help="Tarayıcıyı görsel olarak aç (Varsayılan: headless)")
    parser.add_argument("--output", "-o", default=DEFAULT_OUTPUT_XLSX, help="Çıktı Excel dosyası yolu")
    
    args = parser.parse_args()
    config = load_config()
    
    # Determine search queries
    search_queries = []
    if args.query:
        search_queries = [args.query]
    else:
        presets = config.get("category_presets", {})
        if args.category == "all":
            for cat_key, cat_data in presets.items():
                search_queries.extend(cat_data.get("keywords", []))
        else:
            if args.category in presets:
                search_queries.extend(presets[args.category].get("keywords", []))
            else:
                search_queries = [f"{args.category} İncek Ankara"]
                
    # Run scraper
    results = scrape_google_maps_places(
        search_queries=search_queries,
        anchor_lat=args.lat,
        anchor_lon=args.lon,
        max_places_per_query=args.max_places,
        headless=not args.headed
    )
    
    # Export to Excel
    export_to_excel(results, output_path=args.output)
    
    # Auto-generate or update interactive map
    try:
        from modules.map_generator import generate_interactive_map
        all_leads = []
        if os.path.exists(args.output):
            wb_all = openpyxl.load_workbook(args.output, data_only=True)
            if "TOPLANAN_KURUMLAR" in wb_all.sheetnames:
                ws_a = wb_all["TOPLANAN_KURUMLAR"]
                for r in range(2, ws_a.max_row + 1):
                    all_leads.append({
                        "kurum_adi": ws_a.cell(row=r, column=2).value,
                        "ana_kategori": ws_a.cell(row=r, column=3).value,
                        "is_kolu": ws_a.cell(row=r, column=4).value,
                        "adres": ws_a.cell(row=r, column=5).value,
                        "plus_kodu": ws_a.cell(row=r, column=6).value,
                        "telefon": ws_a.cell(row=r, column=7).value,
                        "eposta": ws_a.cell(row=r, column=8).value,
                        "web_sitesi": ws_a.cell(row=r, column=9).value,
                        "koordinatlar": ws_a.cell(row=r, column=10).value,
                        "mesafe_km": ws_a.cell(row=r, column=11).value,
                        "puan": ws_a.cell(row=r, column=12).value,
                        "hbu_model": ws_a.cell(row=r, column=13).value,
                        "oncelik": ws_a.cell(row=r, column=14).value,
                        "maps_url": ws_a.cell(row=r, column=15).value
                    })
        map_out = os.path.join(os.path.dirname(os.path.abspath(args.output)), "INCEK_TICARI_HEDEF_HARITASI.html")
        generate_interactive_map(all_leads, map_out, anchor_lat=args.lat, anchor_lon=args.lon)
    except Exception as e:
        print(f"Harita güncellenirken uyarı: {e}")

if __name__ == "__main__":
    main()
